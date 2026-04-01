"""
Extract GitHub Project (V2) items into an Excel spreadsheet.

Usage:
    run in GitBash
    export GITHUB_TOKEN="ghp_yourtoken"
    source venv/Scripts/activate # activate venv
    
    python github_project_to_excel.py --owner <org_or_user> --project <project_number> [--output output.xlsx] [--user]
    
    python github_project_to_excel.py --owner drew-csci --project 6 --output project_opportunity_20260323.xlsx
    python github_project_to_excel.py --owner drew-csci --project 5 --output project_minecraft_20260323.xlsx
    python github_project_to_excel.py --owner drew-csci --project 7 --output project_discovery_20260323.xlsx
    
Requirements:
    pip install requests openpyxl

The --user flag indicates the project belongs to a user (default is organization).
A GitHub PAT with read:project scope is required.
"""

import argparse
import os
import sys
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GITHUB_GRAPHQL_URL = "https://api.github.com/graphql"

def run_query(query, variables, token):
    headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    resp = requests.post(GITHUB_GRAPHQL_URL, json={"query": query, "variables": variables}, headers=headers)
    if resp.status_code != 200:
        print(f"HTTP Error {resp.status_code}: {resp.text}", file=sys.stderr)
        sys.exit(1)
    data = resp.json()
    if "errors" in data:
        print(f"GraphQL Errors: {data['errors']}", file=sys.stderr)
        sys.exit(1)
    return data


def fetch_project_items(owner, project_number, token, is_user=False):
    owner_type = "user" if is_user else "organization"
    query = """
    query($owner: String!, $number: Int!, $cursor: String) {
      %s(login: $owner) {
        projectV2(number: $number) {
          title
          fields(first: 50) {
            nodes {
              ... on ProjectV2Field { id name }
              ... on ProjectV2SingleSelectField { id name options { id name } }
              ... on ProjectV2IterationField { id name }
            }
          }
          items(first: 100, after: $cursor) {
            pageInfo { hasNextPage endCursor }
            nodes {
              id
              content {
                ... on Issue {
                  title
                  number
                  state
                  body
                  url
                  createdAt
                  updatedAt
                  author { login }
                  assignees(first: 10) { nodes { login } }
                  labels(first: 10) { nodes { name } }
                  milestone { title }
                  repository { nameWithOwner }
                  comments(first: 50) {
                    nodes {
                      author { login }
                      body
                      createdAt
                    }
                  }
                  closingReferencesConnection: timelineItems(first: 50, itemTypes: [CROSS_REFERENCED_EVENT, CONNECTED_EVENT, DISCONNECTED_EVENT]) {
                    nodes {
                      ... on CrossReferencedEvent {
                        source {
                          ... on PullRequest {
                            number
                            title
                            url
                            state
                            author { login }
                          }
                        }
                      }
                    }
                  }
                }
                ... on PullRequest {
                  title
                  number
                  state
                  body
                  url
                  createdAt
                  updatedAt
                  author { login }
                  assignees(first: 10) { nodes { login } }
                  labels(first: 10) { nodes { name } }
                  milestone { title }
                  repository { nameWithOwner }
                  comments(first: 50) {
                    nodes {
                      author { login }
                      body
                      createdAt
                    }
                  }
                  reviews(first: 50) {
                    nodes {
                      author { login }
                      body
                      state
                      createdAt
                      comments(first: 20) {
                        nodes {
                          author { login }
                          body
                          path
                          createdAt
                        }
                      }
                    }
                  }
                  closingIssuesReferences(first: 10) {
                    nodes {
                      number
                      title
                      url
                    }
                  }
                }
                ... on DraftIssue {
                  title
                  body
                  createdAt
                  updatedAt
                  creator { login }
                }
              }
              fieldValues(first: 50) {
                nodes {
                  ... on ProjectV2ItemFieldTextValue { text field { ... on ProjectV2Field { name } } }
                  ... on ProjectV2ItemFieldNumberValue { number field { ... on ProjectV2Field { name } } }
                  ... on ProjectV2ItemFieldDateValue { date field { ... on ProjectV2Field { name } } }
                  ... on ProjectV2ItemFieldSingleSelectValue { name field { ... on ProjectV2SingleSelectField { name } } }
                  ... on ProjectV2ItemFieldIterationValue { title field { ... on ProjectV2IterationField { name } } }
                }
              }
            }
          }
        }
      }
    }
    """ % owner_type

    all_items = []
    custom_field_names = set()
    cursor = None
    project_title = ""

    while True:
        variables = {"owner": owner, "number": project_number, "cursor": cursor}
        data = run_query(query, variables, token)
        project = data["data"][owner_type]["projectV2"]
        project_title = project["title"]
        items_data = project["items"]

        for node in items_data["nodes"]:
            item = parse_item(node)
            all_items.append(item)
            custom_field_names.update(item.get("custom_fields", {}).keys())

        if items_data["pageInfo"]["hasNextPage"]:
            cursor = items_data["pageInfo"]["endCursor"]
        else:
            break

    return project_title, all_items, sorted(custom_field_names)


def format_comments(comments_nodes):
    if not comments_nodes:
        return ""
    parts = []
    for c in comments_nodes:
        author = c.get("author", {}).get("login", "unknown") if c.get("author") else "unknown"
        date = c.get("createdAt", "")[:10]
        body = (c.get("body") or "").strip()
        if body:
            parts.append(f"[{author} {date}] {body}")
    return "\n---\n".join(parts)


def format_reviews(reviews_nodes):
    if not reviews_nodes:
        return ""
    parts = []
    for r in reviews_nodes:
        author = r.get("author", {}).get("login", "unknown") if r.get("author") else "unknown"
        date = r.get("createdAt", "")[:10]
        state = r.get("state", "")
        body = (r.get("body") or "").strip()
        header = f"[{author} {date} — {state}]"
        if body:
            header += f" {body}"
        
        review_comments = r.get("comments", {}).get("nodes", [])
        rc_parts = []
        for rc in review_comments:
            rc_author = rc.get("author", {}).get("login", "unknown") if rc.get("author") else "unknown"
            rc_path = rc.get("path", "")
            rc_body = (rc.get("body") or "").strip()
            if rc_body:
                rc_parts.append(f"  {rc_path}: {rc_body} ({rc_author})")
        
        if rc_parts:
            header += "\n" + "\n".join(rc_parts)
        parts.append(header)
    return "\n---\n".join(parts)


def format_linked_prs(content):
    parts = []
    
    timeline = content.get("closingReferencesConnection", {}).get("nodes", [])
    for event in timeline:
        source = event.get("source")
        if source and source.get("number"):
            author = source.get("author", {}).get("login", "") if source.get("author") else ""
            parts.append(f"PR #{source['number']} ({source.get('state', '')}) - {source.get('title', '')} by {author} — {source.get('url', '')}")
    
    closing = content.get("closingIssuesReferences", {}).get("nodes", [])
    for issue in closing:
        parts.append(f"Closes #{issue['number']} - {issue.get('title', '')} — {issue.get('url', '')}")
    
    return "\n".join(parts)


def parse_item(node):
    content = node.get("content") or {}
    item = {}

    if "author" in content and content["author"]:
        item["author"] = content["author"]["login"]
    elif "creator" in content and content["creator"]:
        item["author"] = content["creator"]["login"]
    else:
        item["author"] = ""

    item["type"] = "DraftIssue" if "creator" in content else ("PullRequest" if content.get("state") and "url" in content and "/pull/" in content.get("url", "") else "Issue")
    item["title"] = content.get("title", "")
    item["number"] = content.get("number", "")
    item["state"] = content.get("state", "")
    item["body"] = (content.get("body") or "").strip()
    item["repository"] = content.get("repository", {}).get("nameWithOwner", "") if content.get("repository") else ""
    item["url"] = content.get("url", "")
    item["created_at"] = content.get("createdAt", "")
    item["updated_at"] = content.get("updatedAt", "")

    assignees = content.get("assignees", {}).get("nodes", [])
    item["assignees"] = ", ".join(a["login"] for a in assignees) if assignees else ""

    labels = content.get("labels", {}).get("nodes", [])
    item["labels"] = ", ".join(l["name"] for l in labels) if labels else ""

    item["milestone"] = content.get("milestone", {}).get("title", "") if content.get("milestone") else ""

    comments_nodes = content.get("comments", {}).get("nodes", [])
    item["comments"] = format_comments(comments_nodes)
    item["comment_count"] = len(comments_nodes)

    reviews_nodes = content.get("reviews", {}).get("nodes", [])
    item["reviews"] = format_reviews(reviews_nodes)

    item["linked_refs"] = format_linked_prs(content)

    custom = {}
    for fv in node.get("fieldValues", {}).get("nodes", []):
        field = fv.get("field")
        if not field:
            continue
        fname = field.get("name", "")
        if not fname or fname in ("Title",):
            continue
        val = fv.get("text") or fv.get("name") or fv.get("title") or fv.get("date")
        if fv.get("number") is not None and val is None:
            val = fv["number"]
        if val is not None:
            custom[fname] = val
    item["custom_fields"] = custom
    return item


def write_excel(project_title, items, custom_fields, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Project Items"

    fixed = [
        ("Author", "author"),
        ("Type", "type"),
        ("Title", "title"),
        ("#", "number"),
        ("State", "state"),
        ("Repository", "repository"),
        ("Assignees", "assignees"),
        ("Labels", "labels"),
        ("Milestone", "milestone"),
        ("Body", "body"),
        ("Comments (#)", "comment_count"),
        ("Comments", "comments"),
        ("Reviews", "reviews"),
        ("Linked PRs / Issues", "linked_refs"),
        ("Created", "created_at"),
        ("Updated", "updated_at"),
        ("URL", "url"),
    ]
    headers = [h for h, _ in fixed] + custom_fields

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2F5496")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    data_font = Font(name="Arial", size=10)
    wrap_align = Alignment(vertical="top", wrap_text=True)
    alt_fill = PatternFill("solid", fgColor="F2F2F2")

    wrap_columns = {"body", "comments", "reviews", "linked_refs"}

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, item in enumerate(items, 2):
        for col_idx, (_, key) in enumerate(fixed, 1):
            val = item.get(key, "")
            if key in ("created_at", "updated_at") and val:
                val = val[:10]
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if key in wrap_columns:
                cell.alignment = wrap_align
            if row_idx % 2 == 0:
                cell.fill = alt_fill

        for cf_idx, cf_name in enumerate(custom_fields):
            col_idx = len(fixed) + cf_idx + 1
            val = item.get("custom_fields", {}).get(cf_name, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            if row_idx % 2 == 0:
                cell.fill = alt_fill

    wide_cols = {"Body": 50, "Comments": 60, "Reviews": 60, "Linked PRs / Issues": 45}
    for col_idx, header in enumerate(headers, 1):
        if header in wide_cols:
            ws.column_dimensions[get_column_letter(col_idx)].width = wide_cols[header]
        else:
            max_len = len(str(header))
            for row in range(2, min(len(items) + 2, 102)):
                val = ws.cell(row=row, column=col_idx).value
                if val:
                    first_line = str(val).split("\n")[0]
                    max_len = max(max_len, min(len(first_line), 50))
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 3

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(output_path)
    print(f"Saved {len(items)} items from '{project_title}' to {output_path}")


def main():
    parser = argparse.ArgumentParser(description="Export GitHub Project V2 items to Excel")
    parser.add_argument("--owner", required=True, help="GitHub org or username that owns the project")
    parser.add_argument("--project", required=True, type=int, help="Project number (visible in the project URL)")
    parser.add_argument("--output", default="github_project_items.xlsx", help="Output Excel file path")
    parser.add_argument("--user", action="store_true", help="Set if the project belongs to a user (not an org)")
    args = parser.parse_args()

    token = os.environ.get("GITHUB_TOKEN")
    if not token:
        print("Error: Set GITHUB_TOKEN environment variable with a PAT that has read:project scope.", file=sys.stderr)
        sys.exit(1)

    project_title, items, custom_fields = fetch_project_items(args.owner, args.project, token, is_user=args.user)
    write_excel(project_title, items, custom_fields, args.output)

if __name__ == "__main__":
    main()